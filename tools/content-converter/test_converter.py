"""Real engine generation/extraction tests. --require-all forbids native-component skips."""
import base64
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from converter import convert, ConversionError, pdf_for_blocks

REQUIRE_ALL='--require-all' in sys.argv
if REQUIRE_ALL:sys.argv.remove('--require-all')
HAS_OFFICE=bool(shutil.which(os.environ.get('CONVERTER_SOFFICE','soffice')))
HAS_OCR=bool(shutil.which(os.environ.get('CONVERTER_TESSERACT','tesseract')))
HAS_POPPLER=bool(shutil.which(os.environ.get('CONVERTER_PDFTOPPM','pdftoppm')))
if REQUIRE_ALL and not all([HAS_OFFICE,HAS_OCR,HAS_POPPLER]):
    raise SystemExit('REQUIRED_CONVERSION_COMPONENT_MISSING')

SAMPLE='# Weekly report\n\n**Revenue** is ready for review.\n\n| Item | Count |\n| --- | --- |\n| Leads | 12 |\n| Calls | 7 |'


def perform(data,fmt,out,**options):
    if isinstance(data,str):data=data.encode('utf8')
    return convert({'content_base64':base64.b64encode(data).decode(), 'input_format':fmt,'output_format':out,**options})


def output(result,index=0):return base64.b64decode(result['outputs'][index]['content_base64'])


class EngineTests(unittest.TestCase):
    def test_generation_formats(self):
        for fmt in ['pdf','docx','xlsx','txt','html']:
            result=perform(SAMPLE,'md',fmt)
            self.assertTrue(result['ok']);self.assertEqual(result['output_format'],fmt)
            data=output(result);self.assertTrue(data)
            if fmt=='pdf':self.assertIn('Weekly report',PdfReader(io.BytesIO(data)).pages[0].extract_text())
            if fmt=='docx':
                doc=Document(io.BytesIO(data));self.assertEqual(doc.tables[0].cell(1,1).text,'12')
                self.assertTrue(any(run.bold and 'Revenue' in run.text for para in doc.paragraphs for run in para.runs))
            if fmt=='xlsx':
                sheet=load_workbook(io.BytesIO(data)).active;self.assertEqual(sheet['B2'].value,12);self.assertIsInstance(sheet['B2'].value,int)
            if fmt=='html':self.assertIn('<strong>Revenue</strong>',data.decode())
            qa=os.environ.get('MULTI_FORMAT_QA_DIR')
            if qa:Path(qa).mkdir(parents=True,exist_ok=True);(Path(qa)/('sample.'+fmt)).write_bytes(data)

    def test_word_pdf_text(self):
        pdf=output(perform(SAMPLE,'md','pdf'))
        word=output(perform(pdf,'pdf','docx',ocr='off'))
        self.assertIn('Weekly report',' '.join(p.text for p in Document(io.BytesIO(word)).paragraphs))
        text=output(perform(word,'docx','txt')).decode();self.assertIn('Revenue',text)
        if HAS_OFFICE:
            rendered=output(perform(word,'docx','pdf'));self.assertIn('Weekly report',PdfReader(io.BytesIO(rendered)).pages[0].extract_text())
            legacy=output(perform(word,'docx','doc'));self.assertTrue(legacy.startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
            self.assertIn('Weekly report',PdfReader(io.BytesIO(output(perform(legacy,'doc','pdf')))).pages[0].extract_text())

    @unittest.skipUnless(HAS_OFFICE,'LibreOffice absent locally; required on VPS')
    def test_office_render(self):
        for fmt in ['docx','xlsx']:
            file=output(perform(SAMPLE,'md',fmt));pdf=output(perform(file,fmt,'pdf'))
            self.assertGreater(len(PdfReader(io.BytesIO(pdf)).pages),0)
            qa=os.environ.get('MULTI_FORMAT_QA_DIR')
            if qa:(Path(qa)/('rendered-'+fmt+'.pdf')).write_bytes(pdf)

    def test_sanitized_html_and_formula_literals(self):
        unsafe='<h1>Report</h1><script>fetch("http://private/")</script><p onclick="bad()">Safe<img src="file:///etc/passwd"></p><iframe src="http://private"></iframe><table><tr><th>Value</th></tr><tr><td>=HYPERLINK("bad")</td></tr></table>'
        rendered=output(perform(unsafe,'html','html')).decode();self.assertNotIn('private',rendered);self.assertNotIn('onclick',rendered);self.assertNotIn('<script',rendered);self.assertNotIn('file:',rendered)
        book=load_workbook(io.BytesIO(output(perform(unsafe,'html','xlsx'))));self.assertEqual(book.active['A2'].data_type,'s')

    def test_formatted_blocks(self):
        data=json.dumps([{'type':'heading','text':'Summary','level':1},{'type':'table','rows':[['Code','Amount'],['0012',4.5]]}])
        book=load_workbook(io.BytesIO(output(perform(data,'formatted','xlsx'))));self.assertEqual(book.active['A2'].value,'0012');self.assertEqual(book.active['B2'].value,4.5)

    def test_pdf_limits_encryption_and_active_actions(self):
        writer=PdfWriter();writer.add_blank_page(width=612,height=792);writer.encrypt('fixture');out=io.BytesIO();writer.write(out)
        with self.assertRaises(ConversionError):perform(out.getvalue(),'pdf','txt')
        writer=PdfWriter()
        for _ in range(21):writer.add_blank_page(width=612,height=792)
        out=io.BytesIO();writer.write(out)
        with self.assertRaises(ConversionError):perform(out.getvalue(),'pdf','png')
        writer=PdfWriter();writer.add_blank_page(width=612,height=792);writer.add_js('app.alert("fixture")');out=io.BytesIO();writer.write(out)
        with self.assertRaises(ConversionError):perform(out.getvalue(),'pdf','docx')
        from pypdf.generic import DictionaryObject,NameObject,ArrayObject,NumberObject
        writer=PdfWriter();writer.add_blank_page(width=612,height=792)
        writer._root_object[NameObject('/OpenAction')]=DictionaryObject({NameObject('/S'):NameObject('/GoTo'),NameObject('/D'):ArrayObject([NumberObject(0),NameObject('/Fit')])})
        out=io.BytesIO();writer.write(out);self.assertTrue(perform(out.getvalue(),'pdf','png')['ok'])
        writer._root_object[NameObject('/OpenAction')][NameObject('/Next')]=DictionaryObject({NameObject('/S'):NameObject('/Launch')})
        out=io.BytesIO();writer.write(out)
        with self.assertRaises(ConversionError):perform(out.getvalue(),'pdf','png')

    def test_office_macros_rejected(self):
        data=output(perform('Hello','txt','docx'));out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as original,zipfile.ZipFile(out,'w') as modified:
            for member in original.infolist():modified.writestr(member,original.read(member.filename))
            modified.writestr('word/vbaProject.bin',b'fixture')
        with self.assertRaises(ConversionError):perform(out.getvalue(),'docx','txt')

    def test_image_pdf(self):
        image=Image.new('RGB',(500,150),'white');ImageDraw.Draw(image).text((15,30),'INVOICE TOTAL 12345',fill='black')
        data=io.BytesIO();image.save(data,'PNG');pdf=output(perform(data.getvalue(),'png','pdf'))
        self.assertEqual(len(PdfReader(io.BytesIO(pdf)).pages),1)
        if HAS_POPPLER:
            pages=perform(pdf,'pdf','png');self.assertTrue(output(pages).startswith(b'\x89PNG'))

    @unittest.skipUnless(HAS_OCR and HAS_POPPLER,'OCR/native rendering absent locally; required on VPS')
    def test_ocr_image_and_scanned_pdf(self):
        image=Image.new('RGB',(1200,240),'white')
        font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',50)
        ImageDraw.Draw(image).text((35,70),'INVOICE TOTAL 12345',fill='black',font=font)
        buf=io.BytesIO();image.save(buf,'PNG');data=buf.getvalue()
        for fmt,content in [('png',data),('pdf',output(perform(data,'png','pdf')))]:
            text=perform(content,fmt,'txt',operation='ocr');self.assertTrue(text['ocr_used']);self.assertIn('TOTAL 12345',output(text).decode().upper())
        word=perform(data,'png','docx',operation='ocr');self.assertIn('TOTAL 12345',' '.join(p.text for p in Document(io.BytesIO(output(word))).paragraphs).upper())
        qa=os.environ.get('MULTI_FORMAT_QA_DIR')
        if qa:(Path(qa)/'scan.png').write_bytes(data)

    def test_unsupported_and_empty_inputs(self):
        with self.assertRaises(ConversionError):perform('','txt','pdf')
        with self.assertRaises(ConversionError):perform('Hello','txt','exe')
        with self.assertRaises(ConversionError):perform('<script>bad()</script>','html','docx')


if __name__=='__main__':unittest.main(verbosity=2)
